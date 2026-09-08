#!/usr/bin/env python3
"""
Reproduces every figure in analysis/F35-BUDGET-TO-AUDIT.md.

Nothing here writes to Neon or to a page. It reads the parquet warehouse and the
FY2027 budget-exhibit cache and prints the tables in the memo, so any figure can
be re-derived rather than trusted.

    python3 analysis/f35_queries.py --section all
    python3 analysis/f35_queries.py --section 3a

Requires duckdb + openpyxl and the warehouse. Set DM_WAREHOUSE if the warehouse
is not at the default path.
"""
from __future__ import annotations
import argparse, collections, os, sys

PROGRAM_CODE = "198"                      # dod_acquisition_program_code -> 'F-35'
CONTRACT_VINTAGE = "2026-08-06"
PRIOR_VINTAGE    = "2026-07-06"
# The appropriation accounts the P-1 and R-1 F-35 lines sit in.
F35_ACCOUNTS = ("017-1506", "057-3010", "017-1319", "057-3600")
# Accounts that carry F-35 contract obligations but no F-35 budget-exhibit line.
F35_OM_ACCOUNTS = ("017-1804", "057-3400", "057-3840")

def warehouse() -> str:
    for c in (os.environ.get("DM_WAREHOUSE"),
              "/Volumes/AI_DATA/data/usaspending/warehouse",
              os.path.expanduser("~/mnt/warehouse")):
        if c and os.path.isdir(c):
            return c
    sys.exit("Cannot locate the warehouse; set DM_WAREHOUSE")

def budget_cache() -> str:
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    p = os.path.join(here, "database/war_budget_cache/FY2027")
    if not os.path.isdir(p):
        sys.exit(f"Budget exhibit cache not found at {p}")
    return p

def con():
    import duckdb
    c = duckdb.connect(); c.execute("SET memory_limit='6GB'")
    return c

W = warehouse()
G  = lambda v=CONTRACT_VINTAGE: f"{W}/contracts/vintage={v}/fy=*/*.parquet"
FC = f"{W}/accounts/file_c_contracts/fiscal_year=*/*.parquet"
FA = f"{W}/accounts/file_a/fiscal_year=*/*.parquet"
P198 = f"dod_acquisition_program_code='{PROGRAM_CODE}'"

def head(t): print(f"\n{'='*78}\n{t}\n{'='*78}")

# ---------------------------------------------------------------- section 1 --
def is_f35_title(t: str) -> bool:
    """F-35 budget lines. 'Joint Strike Missile' is a different program and is
    deliberately not matched."""
    t = (t or "").upper()
    return "F-35" in t or "JOINT STRIKE FIGHTER" in t or t.startswith("JSF")

def _exhibit(fname, sheet):
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(budget_cache(), fname), read_only=True)
    ws = wb[sheet]; hdr=None; out=[]
    for row in ws.iter_rows(values_only=True):
        if hdr is None:
            if row and row[0] == "Account": hdr = list(row)
            continue
        out.append(list(row))
    wb.close(); return hdr, out

def num(v):
    try: return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError): return 0.0

def section_1():
    """Budget baseline. Add rows only: Weapon System Cost, less prior-year
    advance procurement, plus current-year AP. Non-Add AP detail rows would
    double count. Quantity is summed on cost type A only."""
    head("1. Budget — P-1 (procurement), $ thousands")
    _, rows = _exhibit("p1_display.xlsx", "Exhibit P-1")
    COLS = {"fy25_tot":18, "fy26_disc":20, "fy26_pl119":22, "fy26_tot":24,
            "fy27_disc":26, "fy27_mand":28, "fy27_tot":30}
    QCOL = {"fy25_tot":17, "fy26_tot":23, "fy27_tot":29}
    agg = collections.defaultdict(lambda: collections.defaultdict(float))
    for r in rows:
        if not is_f35_title(r[9]): continue
        if r[12] != "Add": continue
        key = (r[0], r[3], r[8], r[9])
        for k, c in COLS.items(): agg[key][k] += num(r[c])
        if r[10] == "A":
            for k, c in QCOL.items(): agg[key][k+"_q"] += num(r[c])
    tot = collections.defaultdict(float)
    print(f"{'Acct':<7}{'BLI':<8}{'Title':<26}{'FY25':>11}{'q':>4}{'FY26':>11}{'q':>4}{'FY27':>11}{'q':>4}")
    for k in sorted(agg):
        v = agg[k]
        print(f"{k[0]:<7}{k[2]:<8}{k[3][:25]:<26}{v['fy25_tot']:>11,.0f}{v['fy25_tot_q']:>4.0f}"
              f"{v['fy26_tot']:>11,.0f}{v['fy26_tot_q']:>4.0f}{v['fy27_tot']:>11,.0f}{v['fy27_tot_q']:>4.0f}")
        for f in v: tot[f] += v[f]
    print(f"{'TOTAL':<41}{tot['fy25_tot']:>11,.0f}{tot['fy25_tot_q']:>4.0f}"
          f"{tot['fy26_tot']:>11,.0f}{tot['fy26_tot_q']:>4.0f}{tot['fy27_tot']:>11,.0f}{tot['fy27_tot_q']:>4.0f}")

    head("1. Budget — R-1 (RDT&E), $ thousands")
    _, rrows = _exhibit("r1_display.xlsx", "Exhibit R-1")
    rtot = collections.defaultdict(float)
    for r in rrows:
        if not is_f35_title(r[7]): continue
        d = dict(fy25_tot=num(r[11]), fy26_disc=num(r[12]), fy26_pl119=num(r[13]),
                 fy26_tot=num(r[14]), fy27_disc=num(r[15]), fy27_mand=num(r[16]), fy27_tot=num(r[17]))
        print(f"{r[0]:<7}{r[6]:<11}{str(r[7])[:26]:<28}{d['fy25_tot']:>11,.0f}{d['fy26_tot']:>11,.0f}{d['fy27_tot']:>11,.0f}")
        for k, v in d.items(): rtot[k] += v
    print(f"{'TOTAL R-1':<46}{rtot['fy25_tot']:>11,.0f}{rtot['fy26_tot']:>11,.0f}{rtot['fy27_tot']:>11,.0f}")

    head("1a. Discretionary vs mandatory — the FY2027 growth")
    for lbl, a, b in (("Discretionary", tot['fy26_disc']+rtot['fy26_disc'], tot['fy27_disc']+rtot['fy27_disc']),
                      ("Mandatory/PL119-21", tot['fy26_pl119']+rtot['fy26_pl119'], tot['fy27_mand']+rtot['fy27_mand']),
                      ("TOTAL", tot['fy26_tot']+rtot['fy26_tot'], tot['fy27_tot']+rtot['fy27_tot'])):
        chg = (b/a - 1)*100 if a else float("nan")
        print(f"  {lbl:<22} FY26 ${a/1e6:7.3f}B  FY27 ${b/1e6:7.3f}B  {chg:+7.1f}%")
    m = tot['fy27_mand']+rtot['fy27_mand']; t = tot['fy27_tot']+rtot['fy27_tot']
    print(f"  mandatory share of FY2027 request: {m/t*100:.1f}%")

    head("1b. O-1 scan — F-35 lines in the operation and maintenance exhibit")
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(budget_cache(), "o1_display.xlsx"), read_only=True)
    ws = wb[wb.sheetnames[0]]; n = hits = 0; started = False
    for row in ws.iter_rows(values_only=True):
        if not started:
            started = bool(row and row[0] == "Account"); continue
        n += 1
        if is_f35_title(" ".join(str(x) for x in row if x)): hits += 1
    wb.close()
    print(f"  O-1 rows scanned: {n}   rows whose text matches an F-35 line: {hits}")
    print("  (an empty result means the filter matched nothing in this exhibit)")

# ---------------------------------------------------------------- section 2 --
def section_2():
    c = con()
    head("2. Execution — program-198 obligations by fiscal year")
    for fy, a, aw, b in c.execute(f"""SELECT fy, count(*), count(DISTINCT award_id_piid),
      sum(federal_action_obligation)/1e9 FROM read_parquet('{G()}') WHERE {P198} GROUP BY 1 ORDER BY 1""").fetchall():
        ptd = "  (period-to-date)" if fy >= 2026 else ""
        print(f"  FY{fy}  actions {a:>6,}  awards {aw:>5,}  ${b:7.2f}B{ptd}")

    head("2. FY2025 concentration")
    tot = c.execute(f"SELECT sum(federal_action_obligation) FROM read_parquet('{G()}') WHERE {P198} AND fy=2025").fetchone()[0]
    top5 = c.execute(f"""SELECT sum(ob) FROM (SELECT federal_action_obligation ob FROM read_parquet('{G()}')
      WHERE {P198} AND fy=2025 ORDER BY 1 DESC LIMIT 5)""").fetchone()[0]
    piid = c.execute(f"""SELECT sum(federal_action_obligation), count(*) FROM read_parquet('{G()}')
      WHERE {P198} AND fy=2025 AND award_id_piid='N0001923C0003'""").fetchone()
    print(f"  FY2025 total          ${tot/1e9:8.4f}B")
    print(f"  top 5 actions         ${top5/1e9:8.4f}B  ({top5/tot*100:.1f}%)")
    print(f"  PIID N0001923C0003    ${piid[0]/1e9:8.4f}B  ({piid[0]/tot*100:.1f}%, {piid[1]} actions)")
    print("  largest actions:")
    for r in c.execute(f"""SELECT action_date, modification_number, federal_action_obligation/1e9,
      left(coalesce(transaction_description,''),58) FROM read_parquet('{G()}')
      WHERE {P198} AND fy=2025 ORDER BY 3 DESC LIMIT 4""").fetchall():
        print(f"    {r[0]}  {r[1]:<8} ${r[2]:6.3f}B  {r[3]}")

    head("2. Year-end timing — share of obligations in August + September")
    for lbl, where in (("F-35", f"WHERE {P198}"), ("DoD-wide", "")):
        print(f"  {lbl}")
        for fy, t, l in c.execute(f"""SELECT fy, sum(federal_action_obligation)/1e9,
          sum(CASE WHEN month(action_date) IN (8,9) THEN federal_action_obligation ELSE 0 END)/1e9
          FROM read_parquet('{G()}') {where} GROUP BY 1 ORDER BY 1""").fetchall():
            ptd = "  (PTD)" if fy >= 2026 else ""
            print(f"    FY{fy}  ${t:8.2f}B   Aug+Sep ${l:8.2f}B  ({l/t*100:5.1f}%){ptd}")

    head("2. FY2025 composition")
    for lbl, col in (("recipient (parent)","recipient_parent_name"), ("awarding office","awarding_office_name"),
                     ("extent competed","extent_competed"), ("pricing","type_of_contract_pricing")):
        print(f"  -- {lbl}")
        for r in c.execute(f"""SELECT {col}, count(*), sum(federal_action_obligation)/1e9
          FROM read_parquet('{G()}') WHERE {P198} AND fy=2025 GROUP BY 1 ORDER BY 3 DESC LIMIT 4""").fetchall():
            print(f"     {str(r[0])[:44]:<46}{r[1]:>5} actions  ${r[2]:8.3f}B  ({r[2]*1e9/tot*100:5.1f}%)")

# --------------------------------------------------------------- section 3a --
def section_3a():
    c = con()
    head("3a. Account traceability of program-198 obligations")
    NOACCT = "(treasury_accounts_funding_this_award IS NULL OR treasury_accounts_funding_this_award='')"
    for lbl, where in (("F-35 (program 198)", f"WHERE {P198}"), ("DoD-wide", "")):
        print(f"  {lbl}")
        for fy, t, u in c.execute(f"""SELECT fy, sum(federal_action_obligation)/1e9,
          sum(CASE WHEN {NOACCT} THEN federal_action_obligation ELSE 0 END)/1e9
          FROM read_parquet('{G()}') {where} GROUP BY 1 ORDER BY 1""").fetchall():
            ptd = "  (PTD)" if fy >= 2026 else ""
            print(f"    FY{fy}  ${t:8.2f}B   no account named ${u:8.2f}B  ({u/t*100:5.1f}%){ptd}")

    head("3a. Lag test — the same fiscal years at two warehouse vintages")
    print("  If linkage were merely late, a later vintage would repair earlier years.")
    for v in (PRIOR_VINTAGE, CONTRACT_VINTAGE):
        print(f"  -- vintage {v}")
        for fy, t, u, n in c.execute(f"""SELECT fy, sum(federal_action_obligation)/1e9,
          sum(CASE WHEN {NOACCT} THEN federal_action_obligation ELSE 0 END)/1e9, count(*)
          FROM read_parquet('{G(v)}') WHERE {P198} GROUP BY 1 ORDER BY 1""").fetchall():
            print(f"     FY{fy}  ${t:7.2f}B  untraced {u/t*100:5.1f}%  actions {n:,}")

# --------------------------------------------------------------- section 3b --
def section_3b():
    c = con()
    head("3b. Program-198 obligations by exact account combination (FY2022-23)")
    print("  The obligation is NOT apportioned across the accounts named. A single")
    print("  action funded by two appropriations states no split. Do not sum by account.")
    rows = c.execute(f"""SELECT federal_accounts_funding_this_award, federal_action_obligation
      FROM read_parquet('{G()}') WHERE {P198} AND fy IN (2022,2023)
      AND federal_accounts_funding_this_award IS NOT NULL AND federal_accounts_funding_this_award<>''""").fetchall()
    agg = collections.defaultdict(float); cnt = collections.Counter()
    for fa, ob in rows:
        k = ";".join(sorted({a.strip() for a in fa.split(";") if a.strip()}))
        agg[k] += ob or 0; cnt[k] += 1
    for k, v in sorted(agg.items(), key=lambda x: -x[1])[:6]:
        print(f"    ${v/1e9:8.3f}B  ({cnt[k]:>4} actions)  {k}")
    print("\n  011-8242 = Advances, Foreign Military Sales, Funds Appropriated to the")
    print("  President. Agency 011 is out of Department scope per SCOPE-01, but appears")
    print("  commingled with Department accounts on F-35 contract actions.")

# --------------------------------------------------------------- section 3c --
def section_3c():
    c = con()
    head("3c. File C — account-linked contract obligations, F-35 appropriations")
    accts = "','".join(F35_ACCOUNTS)
    for r in c.execute(f"""SELECT fiscal_year, federal_account_symbol, count(*), sum(transaction_obligated_amount)/1e9
      FROM read_parquet('{FC}', hive_partitioning=true)
      WHERE federal_account_symbol IN ('{accts}') GROUP BY 1,2 ORDER BY 1,2""").fetchall():
        print(f"    FY{r[0]}  {r[1]}  rows {r[2]:>9,}  ${r[3]:9.2f}B")
    print("  A fiscal year absent from this listing matched no rows in this warehouse cut.")

    head("3c. File C by PIID — the largest FY2025 F-35 contracts")
    for p in ("N0001923C0003","N0001924C0039","N0001925C0070","N0001923C0030","N0001920C0032"):
        r = c.execute(f"""SELECT count(*), coalesce(sum(transaction_obligated_amount),0)/1e9, min(fiscal_year), max(fiscal_year)
          FROM read_parquet('{FC}', hive_partitioning=true) WHERE award_id_piid='{p}'""").fetchone()
        print(f"    {p}: rows={r[0]:<4} ${r[1]:+.3f}B  fy {r[2]}..{r[3]}")

    head("3c. Department-wide linkage — reproduces the /reconciliation figure")
    fc = dict(c.execute(f"SELECT fiscal_year, sum(transaction_obligated_amount) FROM read_parquet('{FC}', hive_partitioning=true) GROUP BY 1").fetchall())
    aw = dict(c.execute(f"SELECT fy, sum(federal_action_obligation) FROM read_parquet('{G()}') GROUP BY 1").fetchall())
    sp = dict(c.execute(f"SELECT fiscal_year, max(submission_period) FROM read_parquet('{FC}', hive_partitioning=true) GROUP BY 1").fetchall())
    for fy in sorted(aw):
        print(f"    FY{fy}  File C ${fc.get(fy,0)/1e9:7.1f}B / awards ${aw[fy]/1e9:7.1f}B = {fc.get(fy,0)/aw[fy]*100:5.1f}%   latest {sp.get(fy)}")

# --------------------------------------------------------------- section 3d --
def section_3d():
    c = con()
    head("3d. File C linkage by awarding sub-agency, FY2021 vs FY2025")
    d = collections.defaultdict(dict)
    for nm, fy, b, n in c.execute(f"""SELECT awarding_subagency_name, fiscal_year,
      sum(transaction_obligated_amount)/1e9, count(*) FROM read_parquet('{FC}', hive_partitioning=true)
      WHERE fiscal_year IN (2021,2025) GROUP BY 1,2""").fetchall():
        d[nm][fy] = (b, n)
    for nm, v in sorted(d.items(), key=lambda x: -(x[1].get(2021,(0,0))[0]))[:10]:
        a, e = v.get(2021,(0,0)), v.get(2025,(0,0))
        chg = (e[0]/a[0]-1)*100 if a[0] else float("nan")
        print(f"    {str(nm)[:42]:<44} FY21 ${a[0]:7.2f}B  FY25 ${e[0]:7.2f}B  {chg:+7.0f}%")
    print("  Not uniform: a warehouse-wide extract fault would not leave components rising.")

# --------------------------------------------------------------- section 3e --
def section_3e():
    c = con()
    head("3e. File A — the F-35 appropriation accounts have no program dimension")
    accts = "','".join(F35_ACCOUNTS + F35_OM_ACCOUNTS)
    for r in c.execute(f"""SELECT federal_account_symbol, federal_account_name, max(submission_period),
      sum(total_budgetary_resources)/1e9, sum(obligations_incurred)/1e9, sum(gross_outlay_amount)/1e9
      FROM read_parquet('{FA}', hive_partitioning=true)
      WHERE fiscal_year=2025 AND federal_account_symbol IN ('{accts}')
      GROUP BY 1,2 ORDER BY 4 DESC""").fetchall():
        print(f"    {r[0]}  {str(r[1])[:40]:<42}{r[2]}  resources ${r[3]:7.2f}B  oblig ${r[4]:7.2f}B  outlay ${r[5]:7.2f}B")
    print("  These are account totals. Nothing in File A says which part is F-35.")

SECTIONS = {"1":section_1, "2":section_2, "3a":section_3a, "3b":section_3b,
            "3c":section_3c, "3d":section_3d, "3e":section_3e}

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--section", default="all", help="all | " + " | ".join(SECTIONS))
    a = ap.parse_args()
    print(f"warehouse: {W}\ncontract vintage: {CONTRACT_VINTAGE} (prior: {PRIOR_VINTAGE})")
    for k, fn in SECTIONS.items():
        if a.section in ("all", k): fn()
